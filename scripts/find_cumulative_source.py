#!/usr/bin/env python3
"""Find the GOE-level cumulative KPI table source, not a unit-level table.

The golden GOE-164 report shows the intended object under 1.3 GOE-level
Outcomes: one GOE-wide table headed "Cumulative Outcome" with columns
Sr. No / Improvement Areas / Before / Target / After / % Improvement.
This finder deliberately scores that structure above generic files containing
words like "cumulative".
"""
import sys, subprocess, re
from pathlib import Path

root = Path(sys.argv[1])
if not root.exists():
    sys.exit(1)

candidates=[]
for p in root.rglob('*'):
    if not p.is_file():
        continue
    ext=p.suffix.lower()
    if ext not in {'.pdf','.png','.jpg','.jpeg','.docx','.xlsx','.xls','.pptx'}:
        continue
    name=p.name.lower()
    score=0
    text=''
    if 'cumulative' in name: score += 80
    if 'overall' in name: score += 30
    if 'outcome' in name: score += 50
    if 'goe' in name: score += 20
    if 'table' in name: score += 10
    if ext == '.pdf':
        try:
            r=subprocess.run(['pdftotext','-f','1','-l','10',str(p),'-'],capture_output=True,text=True,timeout=20)
            text=r.stdout.lower()
        except Exception:
            pass
    elif ext == '.docx':
        try:
            from docx import Document
            d=Document(str(p))
            chunks=[]
            chunks.extend(par.text for par in d.paragraphs)
            for tbl in d.tables:
                for row in tbl.rows:
                    chunks.append(' | '.join(cell.text for cell in row.cells))
            text='\n'.join(chunks).lower()
        except Exception:
            pass
    elif ext in {'.xlsx','.xls'}:
        try:
            from openpyxl import load_workbook
            wb=load_workbook(str(p), read_only=True, data_only=True)
            chunks=[]
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    chunks.append(' | '.join('' if v is None else str(v) for v in row))
            text='\n'.join(chunks).lower()
        except Exception:
            pass
    elif ext == '.pptx':
        try:
            from pptx import Presentation
            prs=Presentation(str(p)); chunks=[]
            for slide in prs.slides:
                for sh in slide.shapes:
                    if hasattr(sh,'text'): chunks.append(sh.text)
                    if getattr(sh,'has_table',False):
                        for row in sh.table.rows: chunks.append(' | '.join(c.text for c in row.cells))
            text='\n'.join(chunks).lower()
        except Exception:
            pass
    # Exact golden-table structure gets the strongest score.
    if 'cumulative outcome' in text: score += 250
    if 'improvement areas' in text: score += 120
    if 'before' in text and 'target' in text and 'after' in text: score += 100
    if '% improvement' in text or '% improvement' in name: score += 100
    if 'overall outcome (for all units' in text: score += 120
    if 'goe-level outcomes' in text: score += 100
    # A unit KPI table can also contain the same columns, but a filename with
    # a unit name should lose to an explicit GOE/cumulative source.
    if any(k in name for k in ['kpi','niraj','phinix','skipper','reliance','window','dhanashree','phychem','aurumine','monk','reve']):
        score -= 80
    if score:
        candidates.append((score,str(p)))

if not candidates:
    sys.exit(1)
candidates.sort(key=lambda x:(-x[0],x[1].lower()))
print(candidates[0][1])
