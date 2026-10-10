#!/usr/bin/env python3
"""Render the GOE-wide cumulative outcome source into goe_cumulative-*.png.
Supports PDF/images directly and extracts the matching table from DOCX/XLSX.
The key is to preserve the *actual cumulative table data* rather than merely
looking for a filename containing 'cumulative'."""
import sys, re, zipfile, subprocess, os
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont, ImageChops

HEADERS = ['sr. no', 'improvement areas', 'before', 'target', 'after', '% improvement']

def norm(s): return re.sub(r'\s+', ' ', str(s or '').strip().lower())

def trim_whitespace(im, pad=24, tol=12):
    """Crop away blank margin around real content. When a cumulative-outcome
    source is a PDF, pdftoppm rasterizes the WHOLE page -- if the table only
    occupies part of that page, the extra white canvas gets embedded in the
    report right along with it, making the table look small no matter how
    the docx sizes the picture frame. Trim to the content bbox (plus a small
    pad) so the saved PNG's canvas *is* the table."""
    rgb = im.convert('RGB')
    bg = Image.new('RGB', rgb.size, (255, 255, 255))
    diff = ImageChops.difference(rgb, bg)
    diff = ImageChops.add(diff, diff, 2.0, -tol)
    bbox = diff.getbbox()
    if not bbox:
        return im
    l, t, r, b = bbox
    l = max(0, l - pad); t = max(0, t - pad)
    r = min(im.width, r + pad); b = min(im.height, b + pad)
    return im.crop((l, t, r, b))

def font(size, bold=False):
    candidates = [
        '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        'C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf',
    ]
    for p in candidates:
        if os.path.exists(p): return ImageFont.truetype(p, size)
    return ImageFont.load_default()

def render_table(rows, out):
    rows=[[str(c or '').strip() for c in r] for r in rows if any(str(c or '').strip() for c in r)]
    if not rows: raise RuntimeError('No rows found in cumulative table')
    # Find the header row; tolerate a title row above it.
    hi=None
    for i,r in enumerate(rows):
        joined=' | '.join(norm(x) for x in r)
        hits=sum(h in joined for h in HEADERS[1:])
        if hits>=4: hi=i; break
    if hi is None: raise RuntimeError('Could not identify cumulative table header row')
    title='Cumulative Outcome'
    data=rows[hi:]
    # Normalize to 6 columns. If source has extra blank cells, trim them.
    data=[(r+['']*6)[:6] for r in data]
    # Drop an accidental repeated title row.
    data=[r for r in data if norm(r[0])!='cumulative outcome']
    widths=[110,390,170,170,170,180]
    scale=2
    xpad=8; ypad=7
    fs=20; fsb=20
    f=font(fs); fb=font(fsb,True)
    line_h=30
    # Calculate wrapped row heights.
    def wrap(text, max_px, fnt):
        words=str(text).split(); lines=[]; cur=''
        for w in words:
            test=(cur+' '+w).strip()
            if ImageDraw.Draw(Image.new('RGB',(1,1))).textlength(test,font=fnt)<=max_px:
                cur=test
            else:
                if cur: lines.append(cur)
                cur=w
        if cur or not lines: lines.append(cur)
        return lines
    title_h=58
    header_h=58
    row_heights=[]
    for r in data:
        maxlines=max(len(wrap(c,widths[j]-2*xpad,fb if j in (0,1) else f)) for j,c in enumerate(r))
        row_heights.append(max(42,maxlines*line_h+2*ypad))
    W=sum(widths)+2
    H=title_h+header_h+sum(row_heights)+2
    im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    # Title row peach like the reference.
    d.rectangle([1,1,W-2,title_h], fill=(246,204,177), outline='black', width=2)
    tw=d.textbbox((0,0),title,font=fb); d.text(((W-(tw[2]-tw[0]))/2,16),title,font=fb,fill='black')
    headers=['Sr. No','Improvement Areas','Before','Target','After','% Improvement']
    y=title_h
    x=1
    for j,h in enumerate(headers):
        d.rectangle([x,y,x+widths[j],y+header_h],fill=(246,204,177),outline='black',width=1)
        lines=wrap(h,widths[j]-2*xpad,fb); yy=y+(header_h-len(lines)*line_h)/2
        for line in lines:
            bb=d.textbbox((0,0),line,font=fb); d.text((x+(widths[j]-(bb[2]-bb[0]))/2,yy),line,font=fb,fill='black'); yy+=line_h
        x+=widths[j]
    y+=header_h
    for ri,r in enumerate(data):
        x=1
        rh=row_heights[ri]
        for j,c in enumerate(r):
            d.rectangle([x,y,x+widths[j],y+rh],fill='white',outline='black',width=1)
            fnt=fb if (j==1 or norm(c)=='quality') else f
            lines=wrap(c,widths[j]-2*xpad,fnt)
            yy=y+(rh-len(lines)*line_h)/2
            for line in lines:
                bb=d.textbbox((0,0),line,font=fnt)
                if j in (0,2,3,4,5): xx=x+(widths[j]-(bb[2]-bb[0]))/2
                else: xx=x+xpad
                d.text((xx,yy),line,font=fnt,fill='black'); yy+=line_h
            x+=widths[j]
        y+=rh
    im.save(out)

def docx_rows(p):
    import xml.etree.ElementTree as ET
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(p) as z:
        root=ET.fromstring(z.read('word/document.xml'))
        for tbl in root.findall('.//w:tbl',ns):
            rows=[]
            for tr in tbl.findall('./w:tr',ns):
                row=[]
                for tc in tr.findall('./w:tc',ns):
                    row.append(' '.join(t.text or '' for t in tc.findall('.//w:t',ns)))
                rows.append(row)
            joined=' '.join(norm(c) for r in rows for c in r)
            if 'improvement areas' in joined and 'before' in joined and 'target' in joined and 'after' in joined:
                return rows
    raise RuntimeError('No cumulative-style table found in DOCX')

def xlsx_rows(p):
    from openpyxl import load_workbook
    wb=load_workbook(p,read_only=True,data_only=True)
    for ws in wb.worksheets:
        rows=[[str(c.value or '') for c in row] for row in ws.iter_rows()]
        joined=' '.join(norm(c) for r in rows for c in r)
        if 'improvement areas' in joined and 'before' in joined and 'target' in joined and 'after' in joined:
            return rows
    raise RuntimeError('No cumulative-style table found in XLSX')

def main():
    if len(sys.argv)<3: raise SystemExit('usage: render_cumulative_source.py SOURCE OUT_DIR')
    src=Path(sys.argv[1]); out=Path(sys.argv[2]); out.mkdir(parents=True,exist_ok=True)
    for old in out.glob('goe_cumulative-*.png'): old.unlink()
    ext=src.suffix.lower()
    if ext=='.pdf':
        subprocess.run(['pdftoppm','-png','-r','220',str(src),str(out/'goe_cumulative')],check=True)
        # Trim each rasterized page down to its actual content -- pdftoppm
        # renders the full page, and the source page is often much bigger
        # than the table on it.
        for i,p in enumerate(sorted(out.glob('goe_cumulative-*.png')),1):
            im=trim_whitespace(Image.open(p))
            im.save(out/f'goe_cumulative-{i}.png')
            if p != out/f'goe_cumulative-{i}.png': p.unlink()
    elif ext in {'.png','.jpg','.jpeg'}:
        trim_whitespace(Image.open(src).convert('RGB')).save(out/'goe_cumulative-1.png')
    elif ext=='.docx': render_table(docx_rows(src),out/'goe_cumulative-1.png')
    elif ext in {'.xlsx','.xls'}: render_table(xlsx_rows(src),out/'goe_cumulative-1.png')
    else: raise RuntimeError(f'Unsupported cumulative source type: {src.suffix}')
    imgs=sorted(out.glob('goe_cumulative-*.png'))
    if not imgs: raise RuntimeError('Cumulative source produced no images')
    print('Cumulative outcome pages rendered:',len(imgs))
    for p in imgs: print(p)

if __name__=='__main__':
    main()
